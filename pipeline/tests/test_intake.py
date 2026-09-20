import importlib.util
import json
import tempfile
import unittest
from unittest.mock import patch
from argparse import Namespace
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


intake = load_module("intake", ROOT / "pipeline" / "src" / "intake.py")
validator = load_module("validate_case", ROOT / "pipeline" / "src" / "validate_case.py")
import sys
sys.modules["validate_case"] = validator
review_gate = load_module("review_gate", ROOT / "pipeline" / "src" / "review_gate.py")
confirm_input = load_module("confirm_stage1_input", ROOT / "pipeline" / "src" / "confirm_stage1_input.py")


class IntakeTests(unittest.TestCase):
    def make_workbook(self, directory: Path) -> Path:
        path = directory / "input.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "基本信息"
        ws["B1"] = "测试单位"
        ws["D1"] = "测试处室"
        ws["A4"] = "（一）负责甲事项；\n（二）参与乙事项。"
        wb.save(path)
        return path

    def args(self, workbook: Path, output_root: Path) -> Namespace:
        return Namespace(
            workbook=str(workbook), case_id="demo", output_root=str(output_root), title=None,
            sheet="基本信息", unit_cell="B1", department_cell="D1", responsibilities_cell="A4"
        )

    def test_split_responsibilities(self):
        parts = intake.split_responsibilities("（一）负责甲；\n（二）参与乙。")
        self.assertEqual(["（一）负责甲；", "（二）参与乙。"], parts)

    def test_build_and_validate_case(self):
        with tempfile.TemporaryDirectory() as temp:
            temp_path = Path(temp)
            workbook = self.make_workbook(temp_path)
            case_dir = intake.build_case(self.args(workbook, temp_path / "cases"))
            data = json.loads((case_dir / "inputs" / "responsibilities.json").read_text(encoding="utf-8"))
            self.assertEqual(2, len(data["responsibilities"]))
            self.assertEqual("测试处室", data["organization"]["department_name"])
            errors, warnings = validator.validate_case(case_dir)
            self.assertEqual([], errors)
            self.assertTrue(any("尚无 Gate 0" in item for item in warnings))

    def test_rerun_is_idempotent_and_input_change_increments_version(self):
        with tempfile.TemporaryDirectory() as temp:
            temp_path = Path(temp)
            workbook = self.make_workbook(temp_path)
            args = self.args(workbook, temp_path / "cases")
            case_dir = intake.build_case(args)
            intake.build_case(args)
            data = json.loads((case_dir / "inputs" / "responsibilities.json").read_text(encoding="utf-8"))
            self.assertEqual(1, data["input_version"])
            wb = Workbook()
            ws = wb.active
            ws.title = "基本信息"
            ws["B1"] = "测试单位"
            ws["D1"] = "测试处室"
            ws["A4"] = "（一）负责甲；\n（二）参与乙；\n（三）指导丙。"
            wb.save(workbook)
            intake.build_case(args)
            data = json.loads((case_dir / "inputs" / "responsibilities.json").read_text(encoding="utf-8"))
            self.assertEqual(2, data["input_version"])
            self.assertEqual(3, len(data["responsibilities"]))
            self.confirm_declaration(case_dir)
            intake.build_case(args)
            declaration = json.loads((case_dir / "inputs" / "stage-1-input-declaration.json").read_text(encoding="utf-8"))
            self.assertEqual("confirmed", declaration["status"])
            self.assertEqual(2, declaration["input_version"])

    def confirm_declaration(self, case_dir: Path):
        args = Namespace(
            case_dir=str(case_dir), source_type="working-sheet-copy", authority_status="not-verified",
            effective_date=None, source_notes="测试来源", scope_mode="all", include=[], exclude=[],
            scope_notes="测试保留全部", special_rule=[], confirmed_by="tester"
        )
        confirm_input.main = confirm_input.main
        # Call the underlying command through a temporary argv to exercise validation.
        with patch("sys.argv", ["confirm_stage1_input.py", str(case_dir), "--source-type", "working-sheet-copy", "--authority-status", "not-verified", "--source-notes", "测试来源", "--scope-mode", "all", "--scope-notes", "测试保留全部", "--confirmed-by", "tester"]):
            self.assertEqual(0, confirm_input.main())

    def test_advance_requires_approved_decision(self):
        with tempfile.TemporaryDirectory() as temp:
            temp_path = Path(temp)
            workbook = self.make_workbook(temp_path)
            case_dir = intake.build_case(self.args(workbook, temp_path / "cases"))
            advance_args = Namespace(case_dir=str(case_dir))
            with self.assertRaisesRegex(ValueError, "尚无 Gate 0"):
                review_gate.advance(advance_args)
            decide_args = Namespace(case_dir=str(case_dir), decision="approved", reviewer="tester", comments="ok")
            with self.assertRaisesRegex(ValueError, "输入声明尚未"):
                review_gate.record_decision(decide_args)
            self.confirm_declaration(case_dir)
            review_gate.record_decision(decide_args)
            review_gate.advance(advance_args)
            manifest = json.loads((case_dir / "case.json").read_text(encoding="utf-8"))
            self.assertEqual("stage-1", manifest["stage"])
            self.assertEqual("drafting", manifest["status"])
            errors, _ = validator.validate_case(case_dir)
            self.assertEqual([], errors)

    def test_declaration_change_invalidates_gate0_decision(self):
        with tempfile.TemporaryDirectory() as temp:
            temp_path = Path(temp)
            workbook = self.make_workbook(temp_path)
            case_dir = intake.build_case(self.args(workbook, temp_path / "cases"))
            self.confirm_declaration(case_dir)
            decide_args = Namespace(case_dir=str(case_dir), decision="approved", reviewer="tester", comments="ok")
            review_gate.record_decision(decide_args)
            declaration_path = case_dir / "inputs" / "stage-1-input-declaration.json"
            declaration = json.loads(declaration_path.read_text(encoding="utf-8"))
            declaration["special_rules"].append("审核后篡改")
            declaration_path.write_text(json.dumps(declaration, ensure_ascii=False, indent=2), encoding="utf-8")
            errors, _ = validator.validate_case(case_dir)
            self.assertTrue(any("输入声明已失效" in item for item in errors))

    def test_stale_decision_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            temp_path = Path(temp)
            workbook = self.make_workbook(temp_path)
            args = self.args(workbook, temp_path / "cases")
            case_dir = intake.build_case(args)
            source = json.loads((case_dir / "inputs" / "responsibilities.json").read_text(encoding="utf-8"))
            decision = {
                "gate": "gate-0", "decision": "approved", "reviewed_input_version": 99,
                "reviewed_responsibilities_sha256": source["source"]["sha256"],
                "reviewer": "tester", "reviewed_at": "2026-01-01T00:00:00Z"
            }
            (case_dir / "reviews" / "gate-0-decision.json").write_text(
                json.dumps(decision, ensure_ascii=False), encoding="utf-8"
            )
            errors, _ = validator.validate_case(case_dir)
            self.assertTrue(any("输入版本已失效" in item for item in errors))


if __name__ == "__main__":
    unittest.main()
