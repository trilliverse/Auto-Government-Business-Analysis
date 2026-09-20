#!/usr/bin/env python3
"""Export approved Stage 2 assessments to the nine-column review workbook."""
from __future__ import annotations
import argparse,json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Border,Font,PatternFill,Side

def main():
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('case_dir');p.add_argument('--output');a=p.parse_args();case=Path(a.case_dir).resolve();m=json.loads((case/'case.json').read_text(encoding='utf-8'))
 if m.get('stage')!='gate-2' or m.get('status')!='approved':raise ValueError('Gate 2 未批准')
 s1=json.loads((case/m['stage_1_approved_artifact']).read_text(encoding='utf-8'));s2=json.loads((case/m['stage_2_approved_artifact']).read_text(encoding='utf-8'));src=json.loads((case/'inputs/responsibilities.json').read_text(encoding='utf-8'));items={x['item_id']:x for x in s1['items']};resp={x['responsibility_id']:x for x in src['responsibilities']};rows=[]
 for x in s2['assessments']:
  i=items[x['item_id']]; systems='；'.join(y['name'] if y.get('number') is None else f"{y['name']}（{y['number']}）" for y in x['systems']); basis=[]
  for e in x['evidence_refs']:basis.append(f"【直接证据】{e.get('filename','')} {e['locator']}：{e['quote']}")
  basis.extend('【分析整理】'+v for v in x['inference_notes']);basis.extend('【人工补充】'+v for v in x['human_additions'])
  rows.append([src['organization']['department_name'],resp[i['source_responsibility_id']]['source_text'],i['item_text'],i['category'],x['digital_status'],x['process_narrative'],'',systems,'\n'.join(basis)])
 out=Path(a.output).resolve() if a.output else case/'deliverables/养护技术服务处-第二部分人工审核稿.xlsx';out.parent.mkdir(parents=True,exist_ok=True);wb=Workbook();ws=wb.active;ws.title='三定方案业务梳理';headers=['处室名称','职责','事项','职能类型','有没有数字化','流程/待确认问题','泳道图','系统（标号）','流程依据'];ws.append(headers)
 for r in rows:ws.append(r)
 # merge A and consecutive B groups
 if len(rows)>1:ws.merge_cells(start_row=2,start_column=1,end_row=len(rows)+1,end_column=1)
 start=2;prev=items[s2['assessments'][0]['item_id']]['source_responsibility_id']
 for row,x in enumerate(s2['assessments'][1:],start=3):
  cur=items[x['item_id']]['source_responsibility_id']
  if cur!=prev:
   if row-1>start:ws.merge_cells(start_row=start,start_column=2,end_row=row-1,end_column=2)
   start=row;prev=cur
 if len(rows)+1>start:ws.merge_cells(start_row=start,start_column=2,end_row=len(rows)+1,end_column=2)
 thin=Side(style='thin',color='000000');border=Border(left=thin,right=thin,top=thin,bottom=thin);head=PatternFill('solid',fgColor='BFBFBF')
 for c in ws[1]:c.font=Font(bold=True);c.fill=head;c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True);c.border=border
 for r in range(2,len(rows)+2):
  for c in range(1,10):ws.cell(r,c).border=border;ws.cell(r,c).alignment=Alignment(horizontal='center' if c in (1,4,5,7) else 'left',vertical='center',wrap_text=True)
  ws.row_dimensions[r].height=96
 widths=[18,50,42,12,14,80,18,35,90]
 for col,w in zip('ABCDEFGHI',widths):ws.column_dimensions[col].width=w
 ws.freeze_panes='A2';wb.save(out);print(json.dumps({'output':str(out),'item_count':len(rows),'g_nonempty':sum(bool(r[6]) for r in rows)},ensure_ascii=False,indent=2));return 0
if __name__=='__main__':raise SystemExit(main())
