#!/usr/bin/env python3
"""Validate approved Stage 2 nine-column workbook semantics and required blank swimlane column."""
from __future__ import annotations
import argparse,json
from pathlib import Path
from openpyxl import load_workbook

def main():
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('case_dir');p.add_argument('workbook');a=p.parse_args();case=Path(a.case_dir).resolve();m=json.loads((case/'case.json').read_text(encoding='utf-8'));s1=json.loads((case/m['stage_1_approved_artifact']).read_text(encoding='utf-8'));s2=json.loads((case/m['stage_2_approved_artifact']).read_text(encoding='utf-8'));src=json.loads((case/'inputs/responsibilities.json').read_text(encoding='utf-8'));items={x['item_id']:x for x in s1['items']};resp={x['responsibility_id']:x for x in src['responsibilities']};ws=load_workbook(Path(a.workbook),data_only=False,read_only=False)['三定方案业务梳理'];errors=[];headers=['处室名称','职责','事项','职能类型','有没有数字化','流程/待确认问题','泳道图','系统（标号）','流程依据']
 if [ws.cell(1,c).value for c in range(1,10)]!=headers:errors.append('表头不匹配')
 if ws.max_row!=21 or ws.max_column!=9:errors.append(f'尺寸错误 {ws.max_row}x{ws.max_column}')
 dep=responsibility=None
 for row,x in enumerate(s2['assessments'],start=2):
  i=items[x['item_id']]
  if ws.cell(row,1).value is not None:dep=ws.cell(row,1).value
  if ws.cell(row,2).value is not None:responsibility=ws.cell(row,2).value
  actual=[dep,responsibility,ws.cell(row,3).value,ws.cell(row,4).value,ws.cell(row,5).value]
  expected=[src['organization']['department_name'],resp[i['source_responsibility_id']]['source_text'],i['item_text'],i['category'],x['digital_status']]
  if actual!=expected:errors.append(f'第{row}行A-E不匹配')
  if ws.cell(row,7).value not in (None,''):errors.append(f'第{row}行G列应为空')
 if ws.freeze_panes!='A2':errors.append('未冻结首行')
 result={'valid':not errors,'rows':ws.max_row,'columns':ws.max_column,'g_nonempty':sum(ws.cell(r,7).value not in(None,'') for r in range(2,ws.max_row+1)),'errors':errors};print(json.dumps(result,ensure_ascii=False,indent=2));return 1 if errors else 0
if __name__=='__main__':raise SystemExit(main())
