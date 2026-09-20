#!/usr/bin/env python3
"""Run final Stage 1 acceptance checks and write a machine-readable report."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""): h.update(chunk)
    return h.hexdigest()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("case_dir")
    args = parser.parse_args()
    case = Path(args.case_dir).resolve()
    checks = []
    def check(name, condition, detail): checks.append({"name": name, "passed": bool(condition), "detail": detail})
    manifest = json.loads((case / "case.json").read_text(encoding="utf-8"))
    source = json.loads((case / "inputs/responsibilities.json").read_text(encoding="utf-8"))
    declaration_path = case / "inputs/stage-1-input-declaration.json"
    declaration = json.loads(declaration_path.read_text(encoding="utf-8"))
    gate0 = json.loads((case / "reviews/gate-0-decision.json").read_text(encoding="utf-8"))
    gate1 = json.loads((case / "reviews/gate-1-decision.json").read_text(encoding="utf-8"))
    approved_path = case / manifest["stage_1_approved_artifact"]
    approved = json.loads(approved_path.read_text(encoding="utf-8"))
    excel_path = case / "deliverables" / "养护技术服务处-三定方案业务事项清单.xlsx"

    check("工作指引不是案件输入", declaration.get("method_guide_required") is False and "工作指引" not in json.dumps(source, ensure_ascii=False), "method_guide_required=false")
    check("输入声明已确认", declaration.get("status") == "confirmed", declaration.get("confirmed_by"))
    check("Gate 0 绑定来源文件", gate0.get("reviewed_responsibilities_sha256") == source["source"]["sha256"], source["source"]["sha256"])
    check("Gate 0 绑定输入声明", gate0.get("input_declaration_sha256") == sha256(declaration_path), sha256(declaration_path))
    check("Stage 1 批准产物不可变哈希", gate1.get("reviewed_draft_sha256") == sha256(approved_path), sha256(approved_path))
    included = [i for i in approved["items"] if i["disposition"] == "include" and i["status"] != "rejected"]
    excluded = [i for i in approved["items"] if i["disposition"] != "include" or i["status"] == "rejected"]
    check("职责数量", len(source["responsibilities"]) == 8, str(len(source["responsibilities"])))
    check("正式事项数量", len(included) == 20, str(len(included)))
    check("排除项审计保留", len(excluded) == 1 and excluded[0]["source_responsibility_id"] == "R008", json.dumps(excluded, ensure_ascii=False))
    check("Excel 存在", excel_path.is_file(), str(excel_path))
    if excel_path.is_file():
        ws = load_workbook(excel_path, data_only=False, read_only=False)["三定方案业务梳理"]
        check("Excel 四列二十项", ws.max_column == 4 and ws.max_row == 21, f"rows={ws.max_row}, cols={ws.max_column}")
        check("Excel 合并范围", "A2:A21" in {str(r) for r in ws.merged_cells.ranges}, str(list(ws.merged_cells.ranges)))
    result = {"case_id": manifest["case_id"], "passed": all(c["passed"] for c in checks), "checks": checks}
    output = case / "reviews" / "stage-1-acceptance.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
