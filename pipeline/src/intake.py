#!/usr/bin/env python3
"""Stage 0 intake: extract organization and responsibility clauses from Excel."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SCHEMA_VERSION = "0.1.0"
CLAUSE_PATTERN = re.compile(r"(?m)(?=（[一二三四五六七八九十百]+）)")
CLAUSE_HEAD_PATTERN = re.compile(r"^（([一二三四五六七八九十百]+)）")


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def split_responsibilities(text: str) -> list[str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        raise ValueError("职责单元格为空")
    parts = [part.strip() for part in CLAUSE_PATTERN.split(normalized) if part.strip()]
    if not parts or any(CLAUSE_HEAD_PATTERN.match(part) is None for part in parts):
        raise ValueError("无法按中文括号条款编号切分职责；请人工确认单元格或编号格式")
    return parts


def read_cell(ws: Any, ref: str, label: str) -> str:
    value = ws[ref].value
    if value is None or not str(value).strip():
        raise ValueError(f"{label}单元格 {ws.title}!{ref} 为空")
    return str(value).strip()


def make_review_markdown(data: dict[str, Any], case_id: str) -> str:
    source = data["source"]
    lines = [
        "# Gate 0｜输入材料与职责原文审核",
        "",
        f"- 案件 ID：`{case_id}`",
        f"- 单位：{data['organization']['unit_name']}",
        f"- 处室：{data['organization']['department_name']}",
        f"- 来源：`{source['locator']}`",
        f"- 输入版本：`{data['input_version']}`",
        f"- 文件 SHA-256：`{source['sha256']}`",
        f"- 职责数量：**{len(data['responsibilities'])}**",
        "",
        "## 审核说明",
        "",
        "请核对单位、处室、职责原文、条款数量与切分边界。此稿只用于 Gate 0；未批准前不得开始职责拆解。",
        "",
    ]
    for item in data["responsibilities"]:
        lines.extend([
            f"## {item['responsibility_id']}｜第 {item['sequence']} 条",
            "",
            item["source_text"],
            "",
            "- [ ] 原文完整",
            "- [ ] 条款边界正确",
            "- [ ] 纳入范围与特殊口径已确认",
            "- 审核意见：",
            "",
        ])
    lines.extend([
        "## Gate 0 决定",
        "",
        "完成核对后，先确认 `inputs/stage-1-input-declaration.json` 中的职责来源权威性、纳入排除范围和特殊口径，再创建 Gate 0 决定。工作指引方法已内置，不是案件附件。",
        "",
    ])
    return "\n".join(lines)


def build_case(args: argparse.Namespace) -> Path:
    workbook = Path(args.workbook).resolve()
    if not workbook.is_file():
        raise FileNotFoundError(f"工作簿不存在：{workbook}")

    wb = load_workbook(workbook, data_only=False, read_only=False)
    if args.sheet not in wb.sheetnames:
        raise ValueError(f"工作表不存在：{args.sheet}；可用：{', '.join(wb.sheetnames)}")
    ws = wb[args.sheet]
    unit_name = read_cell(ws, args.unit_cell, "单位")
    department_name = read_cell(ws, args.department_cell, "处室")
    raw_text = read_cell(ws, args.responsibilities_cell, "职责")
    clauses = split_responsibilities(raw_text)
    file_hash = sha256_file(workbook)

    case_dir = Path(args.output_root).resolve() / args.case_id
    inputs_dir = case_dir / "inputs"
    reviews_dir = case_dir / "reviews"
    (case_dir / "stage-1" / "drafts").mkdir(parents=True, exist_ok=True)
    (case_dir / "stage-1" / "approved").mkdir(parents=True, exist_ok=True)
    inputs_dir.mkdir(parents=True, exist_ok=True)
    reviews_dir.mkdir(parents=True, exist_ok=True)

    now = utc_now()
    existing_manifest = case_dir / "case.json"
    created_at = now
    input_version = 1
    if existing_manifest.exists():
        previous = json.loads(existing_manifest.read_text(encoding="utf-8"))
        created_at = previous.get("created_at", now)
        previous_source_path = inputs_dir / "responsibilities.json"
        if previous_source_path.exists():
            previous_source = json.loads(previous_source_path.read_text(encoding="utf-8"))
            previous_hash = previous_source.get("source", {}).get("sha256")
            previous_version = int(previous_source.get("input_version", 1))
            input_version = previous_version if previous_hash == file_hash else previous_version + 1

    source_locator = f"{workbook.name}#{args.sheet}!{args.responsibilities_cell}"
    responsibilities = [
        {
            "responsibility_id": f"R{index:03d}",
            "sequence": index,
            "source_text": clause,
            "source_locator": source_locator,
        }
        for index, clause in enumerate(clauses, start=1)
    ]
    source_data = {
        "schema_version": SCHEMA_VERSION,
        "case_id": args.case_id,
        "input_version": input_version,
        "organization": {"unit_name": unit_name, "department_name": department_name},
        "source": {
            "type": "xlsx",
            "path": str(workbook),
            "filename": workbook.name,
            "sha256": file_hash,
            "sheet": args.sheet,
            "unit_cell": args.unit_cell,
            "department_cell": args.department_cell,
            "responsibilities_cell": args.responsibilities_cell,
            "locator": source_locator,
        },
        "responsibilities": responsibilities,
        "extracted_at": now,
    }
    manifest = {
        "schema_version": SCHEMA_VERSION,
        "case_id": args.case_id,
        "title": args.title or f"{department_name}职责与数字化流程梳理",
        "organization": {"unit_name": unit_name, "department_name": department_name},
        "stage": "gate-0",
        "status": "awaiting-input-review",
        "input_version": input_version,
        "inputs": {"responsibilities": "inputs/responsibilities.json"},
        "created_at": created_at,
        "updated_at": now,
    }

    (inputs_dir / "responsibilities.json").write_text(
        json.dumps(source_data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    existing_manifest.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    declaration_path = inputs_dir / "stage-1-input-declaration.json"
    declaration_needs_reset = not declaration_path.exists()
    if declaration_path.exists():
        try:
            existing_declaration = json.loads(declaration_path.read_text(encoding="utf-8"))
            declaration_needs_reset = existing_declaration.get("input_version") != input_version
        except (OSError, json.JSONDecodeError):
            declaration_needs_reset = True
    if declaration_needs_reset:
        declaration = {
            "schema_version": SCHEMA_VERSION,
            "case_id": args.case_id,
            "input_version": input_version,
            "responsibility_source": {
                "source_type": "working-sheet-copy",
                "authority_status": "not-verified",
                "effective_date": None,
                "notes": "默认仅确认从工作簿提取；是否为现行正式三定文本需人工确认。"
            },
            "scope": {
                "mode": "all",
                "included_responsibility_ids": [],
                "excluded_responsibility_ids": [],
                "notes": "默认保留全部职责进入草稿；正式纳入或排除范围需人工确认。"
            },
            "special_rules": [],
            "method_guide_required": False,
            "status": "pending-human-confirmation",
            "confirmed_by": None,
            "confirmed_at": None
        }
        declaration_path.write_text(json.dumps(declaration, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (reviews_dir / "gate-0-input-review.md").write_text(
        make_review_markdown(source_data, args.case_id), encoding="utf-8"
    )
    print(json.dumps({
        "case_dir": str(case_dir),
        "unit_name": unit_name,
        "department_name": department_name,
        "responsibility_count": len(responsibilities),
        "input_version": input_version,
        "sha256": file_hash,
        "status": "awaiting-input-review",
    }, ensure_ascii=False, indent=2))
    return case_dir


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--case-id", required=True)
    parser.add_argument("--output-root", default="workcases")
    parser.add_argument("--title")
    parser.add_argument("--sheet", default="基本信息")
    parser.add_argument("--unit-cell", default="B1")
    parser.add_argument("--department-cell", default="D1")
    parser.add_argument("--responsibilities-cell", default="A4")
    return parser.parse_args()


if __name__ == "__main__":
    build_case(parse_args())
