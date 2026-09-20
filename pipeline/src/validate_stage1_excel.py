#!/usr/bin/env python3
"""Validate Stage 1 Excel semantics, merges, and essential presentation rules."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("case_dir")
    parser.add_argument("workbook")
    args = parser.parse_args()
    case_dir = Path(args.case_dir).resolve()
    workbook = Path(args.workbook).resolve()
    manifest = json.loads((case_dir / "case.json").read_text(encoding="utf-8"))
    source = json.loads((case_dir / "inputs" / "responsibilities.json").read_text(encoding="utf-8"))
    artifact = json.loads((case_dir / manifest["stage_1_approved_artifact"]).read_text(encoding="utf-8"))
    source_map = {item["responsibility_id"]: item for item in source["responsibilities"]}
    expected = [item for item in artifact["items"] if item["disposition"] == "include" and item["status"] != "rejected"]
    wb = load_workbook(workbook, data_only=False, read_only=False)
    ws = wb["三定方案业务梳理"]
    errors = []
    if [ws.cell(1, c).value for c in range(1, 5)] != ["处室名称", "职责", "事项", "职能类型"]:
        errors.append("表头不匹配")
    if ws.max_row != len(expected) + 1 or ws.max_column != 4:
        errors.append(f"尺寸不匹配：{ws.max_row}x{ws.max_column}")
    current_department = None
    current_responsibility = None
    for row, item in enumerate(expected, start=2):
        if ws.cell(row, 1).value is not None: current_department = ws.cell(row, 1).value
        if ws.cell(row, 2).value is not None: current_responsibility = ws.cell(row, 2).value
        expected_row = [source["organization"]["department_name"], source_map[item["source_responsibility_id"]]["source_text"], item["item_text"], item["category"]]
        actual_row = [current_department, current_responsibility, ws.cell(row, 3).value, ws.cell(row, 4).value]
        if actual_row != expected_row: errors.append(f"第 {row} 行内容不匹配")
        for col in range(1, 5):
            cell = ws.cell(row, col)
            # openpyxl represents non-anchor cells inside merged ranges as
            # MergedCell and synthesizes edge borders; validate the anchor and
            # every ordinary cell instead of requiring each placeholder to
            # retain the pre-merge style object.
            if isinstance(cell, MergedCell):
                continue
            if cell.border.left.style != "thin" or cell.border.right.style != "thin" or not cell.alignment.wrap_text:
                errors.append(f"第 {row} 行第 {col} 列样式不完整")
    merges = {str(value) for value in ws.merged_cells.ranges}
    expected_a = f"A2:A{len(expected)+1}"
    if expected_a not in merges: errors.append(f"缺少处室合并区域 {expected_a}")
    # Derive expected B merges from consecutive source ids.
    start = 2
    previous = expected[0]["source_responsibility_id"]
    expected_b = set()
    for offset, item in enumerate(expected[1:], start=3):
        if item["source_responsibility_id"] != previous:
            if offset - 1 > start: expected_b.add(f"B{start}:B{offset-1}")
            start = offset
            previous = item["source_responsibility_id"]
    if len(expected) + 1 > start: expected_b.add(f"B{start}:B{len(expected)+1}")
    actual_b = {value for value in merges if value.startswith("B")}
    if actual_b != expected_b: errors.append(f"B 列合并不匹配：expected={sorted(expected_b)} actual={sorted(actual_b)}")
    if ws.freeze_panes != "A2": errors.append("未冻结首行")
    result = {"valid": not errors, "item_count": len(expected), "merged_ranges": sorted(merges), "errors": errors}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
