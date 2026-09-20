#!/usr/bin/env python3
"""Export the immutable approved Stage 1 artifact to a styled four-column workbook."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("case_dir")
    parser.add_argument("--output")
    args = parser.parse_args()
    case_dir = Path(args.case_dir).resolve()
    manifest = json.loads((case_dir / "case.json").read_text(encoding="utf-8"))
    if manifest.get("stage") != "gate-1" or manifest.get("status") != "approved":
        raise ValueError("只有 Gate 1 已批准案件可以导出 Stage 1 Excel")
    artifact_path = case_dir / manifest["stage_1_approved_artifact"]
    artifact = json.loads(artifact_path.read_text(encoding="utf-8"))
    source = json.loads((case_dir / "inputs" / "responsibilities.json").read_text(encoding="utf-8"))
    source_map = {item["responsibility_id"]: item for item in source["responsibilities"]}
    items = [item for item in artifact["items"] if item["disposition"] == "include" and item["status"] != "rejected"]
    if not items:
        raise ValueError("没有可导出的批准事项")

    output = Path(args.output).resolve() if args.output else case_dir / "deliverables" / f"{source['organization']['department_name']}-三定方案业务事项清单.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "三定方案业务梳理"
    headers = ["处室名称", "职责", "事项", "职能类型"]
    ws.append(headers)
    for item in items:
        ws.append([
            source["organization"]["department_name"],
            source_map[item["source_responsibility_id"]]["source_text"],
            item["item_text"],
            item["category"],
        ])

    # Merge A for the department and B for each consecutive responsibility group.
    if len(items) > 1:
        ws.merge_cells(start_row=2, start_column=1, end_row=len(items) + 1, end_column=1)
    group_start = 2
    previous = items[0]["source_responsibility_id"]
    for offset, item in enumerate(items[1:], start=3):
        if item["source_responsibility_id"] != previous:
            if offset - 1 > group_start:
                ws.merge_cells(start_row=group_start, start_column=2, end_row=offset - 1, end_column=2)
            group_start = offset
            previous = item["source_responsibility_id"]
    if len(items) + 1 > group_start:
        ws.merge_cells(start_row=group_start, start_column=2, end_row=len(items) + 1, end_column=2)

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="BFBFBF")
    group_fills = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="E7E6E6")]
    for cell in ws[1]:
        cell.font = Font(bold=True, color="000000")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    source_order = []
    for item in items:
        if item["source_responsibility_id"] not in source_order:
            source_order.append(item["source_responsibility_id"])
    fill_by_source = {source_id: group_fills[index % 2] for index, source_id in enumerate(source_order)}
    for row_index, item in enumerate(items, start=2):
        for column in range(1, 5):
            cell = ws.cell(row_index, column)
            cell.border = border
            cell.fill = fill_by_source[item["source_responsibility_id"]]
            cell.alignment = Alignment(horizontal="center" if column in (1, 4) else "left", vertical="center", wrap_text=True)
        ws.row_dimensions[row_index].height = 42
    ws.row_dimensions[1].height = 28
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 54
    ws.column_dimensions["C"].width = 48
    ws.column_dimensions["D"].width = 12
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{len(items)+1}"
    wb.save(output)
    print(json.dumps({"output": str(output), "department_count": 1, "responsibility_count": len(source_order), "item_count": len(items)}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
