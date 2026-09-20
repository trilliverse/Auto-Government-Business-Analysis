#!/usr/bin/env python3
"""Build a deterministic Stage 2 document catalog and locator-preserving text index."""

from __future__ import annotations

import argparse
import hashlib
import json
import mimetypes
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from pypdf import PdfReader

SUPPORTED = {".pdf", ".xlsx", ".docx", ".md", ".txt"}


def infer_system_metadata(path: Path) -> dict:
    parts = path.parts
    system_dir = next((part for part in parts if re.match(r"^(?:\d|3\+10)", part)), path.parent.name)
    match = re.match(r"^(3\+10|\d+)\s*(.*)$", system_dir)
    number = match.group(1) if match else None
    name = match.group(2).strip() if match else system_dir
    filename = path.name.lower()
    if "需求" in filename: doc_type = "requirements"
    elif "概要" in filename: doc_type = "high-level-design"
    elif "详细" in filename: doc_type = "detailed-design"
    elif "数据库" in filename or "数据设计" in filename: doc_type = "database-design"
    elif "操作" in filename or "维护手册" in filename: doc_type = "operation-manual"
    elif "初步设计" in filename: doc_type = "preliminary-design"
    elif "可行性" in filename or "工可" in filename: doc_type = "feasibility-study"
    elif "测试" in filename: doc_type = "test-material"
    else: doc_type = "other"
    scopes = []
    combined = system_dir + path.name
    for token in ("普通国省干线", "普通干线", "高速公路", "农村公路", "公路水运", "交通流量", "应急", "养护"):
        if token in combined and token not in scopes: scopes.append(token)
    return {"system_number": number, "system_name": name, "system_directory": system_dir, "document_type": doc_type, "scope_tags": scopes}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def chunks_from_pdf(path: Path) -> Iterable[dict]:
    reader = PdfReader(str(path))
    for page_number, page in enumerate(reader.pages, start=1):
        text = normalize(page.extract_text() or "")
        if text:
            yield {"locator": f"page:{page_number}", "text": text, "kind": "page"}


def chunks_from_xlsx(path: Path) -> Iterable[dict]:
    wb = load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            values = [normalize(str(cell.value)) for cell in row if cell.value not in (None, "")]
            if not values:
                continue
            joined = " | ".join(values)
            # Data minimization: contact rows are irrelevant to process evidence
            # and must not enter the searchable index or review report.
            if any(label in joined for label in ("联络人姓名", "联系电话", "电话 |")):
                continue
            if ws.title == "基本信息": role = "responsibility-source"
            elif "流程" in ws.title or "事项" in ws.title: role = "process-source"
            else: role = "business-context"
            yield {"locator": f"sheet:{ws.title}!row:{row[0].row}", "text": joined, "kind": "row", "evidence_role": role}


def chunks_from_docx(path: Path) -> Iterable[dict]:
    # DOCX is a ZIP package. Use stdlib only to avoid another runtime dependency.
    import zipfile
    import xml.etree.ElementTree as ET
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraph = 0
    for node in root.iter(ns + "p"):
        text = normalize("".join(child.text or "" for child in node.iter(ns + "t")))
        if text:
            paragraph += 1
            yield {"locator": f"paragraph:{paragraph}", "text": text, "kind": "paragraph"}


def chunks_from_text(path: Path) -> Iterable[dict]:
    for line_number, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), start=1):
        text = normalize(line)
        if text:
            yield {"locator": f"line:{line_number}", "text": text, "kind": "line"}


def extract(path: Path) -> list[dict]:
    if path.suffix.lower() == ".pdf": return list(chunks_from_pdf(path))
    if path.suffix.lower() == ".xlsx": return list(chunks_from_xlsx(path))
    if path.suffix.lower() == ".docx": return list(chunks_from_docx(path))
    return list(chunks_from_text(path))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("case_dir")
    parser.add_argument("--input", action="append", required=True, help="File or directory; repeatable")
    parser.add_argument("--scope-note", default="用户提供的 Stage 2 分析材料；适用范围待人工确认")
    args = parser.parse_args()
    case_dir = Path(args.case_dir).resolve()
    manifest = json.loads((case_dir / "case.json").read_text(encoding="utf-8"))
    if manifest.get("stage") != "stage-2" or manifest.get("status") != "indexing":
        raise ValueError("案件必须处于 stage-2/indexing")
    paths = []
    for value in args.input:
        target = Path(value).resolve()
        if target.is_file(): paths.append(target)
        elif target.is_dir(): paths.extend(p for p in target.rglob("*") if p.is_file())
        else: raise FileNotFoundError(target)
    paths = sorted({p for p in paths if p.suffix.lower() in SUPPORTED and not p.name.startswith(("~$", ".~", "._"))}, key=lambda p: str(p).lower())
    catalog = []
    all_chunks = []
    for index, path in enumerate(paths, start=1):
        document_id = f"DOC{index:04d}"
        digest = sha256(path)
        try:
            chunks = extract(path)
            extraction_status = "ok"
            extraction_error = None
        except Exception as exc:
            chunks = []
            extraction_status = "failed"
            extraction_error = f"{type(exc).__name__}: {exc}"
        metadata = infer_system_metadata(path)
        catalog.append({
            "document_id": document_id, "path": str(path), "filename": path.name,
            "sha256": digest, "media_type": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
            **metadata,
            "scope_note": args.scope_note, "extraction_status": extraction_status,
            "extraction_error": extraction_error, "chunk_count": len(chunks)
        })
        for chunk_number, chunk in enumerate(chunks, start=1):
            all_chunks.append({
                "chunk_id": f"{document_id}-C{chunk_number:05d}", "document_id": document_id,
                "locator": chunk["locator"], "kind": chunk["kind"], "evidence_role": chunk.get("evidence_role", "reference-method" if path.suffix.lower() == ".pdf" else "business-context"), "text": chunk["text"]
            })
    result = {
        "schema_version": "0.1.0", "case_id": manifest["case_id"],
        "stage_1_artifact_sha256": manifest["stage_1_approved_sha256"],
        "created_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "documents": catalog, "chunks": all_chunks
    }
    output = case_dir / "stage-2" / "index" / "document-index.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(output), "document_count": len(catalog), "chunk_count": len(all_chunks), "failed": sum(d["extraction_status"] == "failed" for d in catalog)}, ensure_ascii=False, indent=2))
    return 1 if any(d["extraction_status"] == "failed" for d in catalog) else 0


if __name__ == "__main__":
    raise SystemExit(main())
