#!/usr/bin/env python3
"""Extract a Stage 1 reference set from an existing reviewed workbook."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

ALLOWED_CATEGORIES = {"督导", "组织", "负责", "参与", "待确认"}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--sheet", default="三定方案业务梳理")
    parser.add_argument("--case-id", default="maintenance-office-demo")
    args = parser.parse_args()

    workbook = Path(args.workbook).resolve()
    wb = load_workbook(workbook, data_only=False, read_only=False)
    if args.sheet not in wb.sheetnames:
        raise ValueError(f"工作表不存在：{args.sheet}")
    ws = wb[args.sheet]
    headers = {str(ws.cell(1, col).value).strip(): col for col in range(1, ws.max_column + 1) if ws.cell(1, col).value}
    required = ["处室名称", "职责", "事项", "职能类型"]
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError("缺少列：" + "、".join(missing))

    current_department = None
    current_responsibility = None
    responsibility_ids: dict[str, str] = {}
    source_item_counts: dict[str, int] = {}
    items = []
    for row in range(2, ws.max_row + 1):
        department = ws.cell(row, headers["处室名称"]).value
        responsibility = ws.cell(row, headers["职责"]).value
        item_text = ws.cell(row, headers["事项"]).value
        category = ws.cell(row, headers["职能类型"]).value
        if department not in (None, ""):
            current_department = str(department).strip()
        if responsibility not in (None, ""):
            current_responsibility = str(responsibility).strip()
        if item_text in (None, ""):
            continue
        if current_department is None or current_responsibility is None:
            raise ValueError(f"第 {row} 行无法向上继承处室或职责")
        category = str(category).strip() if category is not None else ""
        if category not in ALLOWED_CATEGORIES:
            raise ValueError(f"第 {row} 行分类无效：{category}")
        if current_responsibility not in responsibility_ids:
            responsibility_ids[current_responsibility] = f"REF-R{len(responsibility_ids)+1:03d}"
        source_id = responsibility_ids[current_responsibility]
        source_item_counts[source_id] = source_item_counts.get(source_id, 0) + 1
        items.append({
            "reference_item_id": f"REF-I{len(items)+1:03d}",
            "source_responsibility_id": source_id,
            "source_responsibility_text": current_responsibility,
            "source_sequence": list(responsibility_ids.values()).index(source_id) + 1,
            "item_sequence": source_item_counts[source_id],
            "department_name": current_department,
            "item_text": str(item_text).strip(),
            "category": category,
            "source_row": row,
        })

    output = {
        "schema_version": "0.1.0",
        "case_id": args.case_id,
        "purpose": "stage-1-human-reviewed-reference",
        "source": {"path": str(workbook), "filename": workbook.name, "sheet": args.sheet},
        "statistics": {
            "responsibility_count": len(responsibility_ids),
            "item_count": len(items),
            "category_counts": {name: sum(1 for item in items if item["category"] == name) for name in sorted(ALLOWED_CATEGORIES)},
        },
        "items": items,
        "limitations": [
            "该工作簿是人工核验后的后续阶段成果，仅作为 Stage 1 回归参考，不反向替代原始职责。",
            "该参考集已合并方针政策与法律法规事项，并排除了领导交办事项，因此预期为 7 条来源职责、20 条事项。"
        ]
    }
    output_path = Path(args.output).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(output["statistics"], ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
