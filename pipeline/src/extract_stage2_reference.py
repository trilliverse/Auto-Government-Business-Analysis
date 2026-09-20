#!/usr/bin/env python3
"""Extract a Stage 2 reviewed reference from the example nine-column workbook."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

STATUSES = {"有", "无", "待确认", "不纳入本次梳理范畴"}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--sheet", default="三定方案业务梳理")
    parser.add_argument("--stage1-reference", required=True)
    args = parser.parse_args()
    workbook = Path(args.workbook).resolve()
    stage1 = json.loads(Path(args.stage1_reference).resolve().read_text(encoding="utf-8"))
    expected_by_pair = {(i["item_text"], i["category"]): i for i in stage1["items"]}
    wb = load_workbook(workbook, data_only=False, read_only=False)
    ws = wb[args.sheet]
    headers = {str(ws.cell(1, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    required = ["事项", "职能类型", "有没有数字化", "流程/待确认问题", "系统（标号）", "流程依据"]
    missing = [name for name in required if name not in headers]
    if missing: raise ValueError("缺少列：" + "、".join(missing))
    items = []
    for row in range(2, ws.max_row + 1):
        item_text = ws.cell(row, headers["事项"]).value
        category = ws.cell(row, headers["职能类型"]).value
        if item_text in (None, ""): continue
        key = (str(item_text).strip(), str(category).strip())
        stage1_item = expected_by_pair.get(key)
        if stage1_item is None: raise ValueError(f"第 {row} 行无法映射 Stage 1 Gold：{key}")
        digital_status = str(ws.cell(row, headers["有没有数字化"]).value or "").strip()
        if digital_status not in STATUSES: raise ValueError(f"第 {row} 行数字化状态无效：{digital_status}")
        narrative = str(ws.cell(row, headers["流程/待确认问题"]).value or "").strip()
        systems_text = str(ws.cell(row, headers["系统（标号）"]).value or "").strip()
        basis = str(ws.cell(row, headers["流程依据"]).value or "").strip()
        systems = [part.strip() for part in systems_text.split("；") if part.strip()]
        items.append({
            "item_id": stage1_item["reference_item_id"],
            "item_text": stage1_item["item_text"],
            "category": stage1_item["category"],
            "source_row": row,
            "digital_status": digital_status,
            "process_narrative": narrative,
            "systems": systems,
            "evidence_narrative": basis,
            "contains_unresolved": "待确认" in narrative,
        })
    result = {
        "schema_version": "0.1.0", "purpose": "stage-2-human-reviewed-reference",
        "source": {"path": str(workbook), "filename": workbook.name, "sheet": args.sheet},
        "statistics": {
            "item_count": len(items),
            "digital_status_counts": {value: sum(i["digital_status"] == value for i in items) for value in sorted(STATUSES)},
            "items_with_unresolved": sum(i["contains_unresolved"] for i in items),
            "items_with_systems": sum(bool(i["systems"]) for i in items),
            "items_with_evidence_narrative": sum(bool(i["evidence_narrative"]) for i in items)
        },
        "items": items,
        "limitations": [
            "该文件是人工核验结果，用于开发回归，不等同于所有结论均已被原始系统资料独立证实。",
            "流程依据列是人工成果中的叙述性依据；后续生产流程需拆成文档ID、定位、短引文和证据强度。"
        ]
    }
    output = Path(args.output).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result["statistics"], ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
